from pathlib import Path
import json,html,shutil,re
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent;ROOT=HERE.parent.parent
BASE='Thuc_hanh_CTDL_GT_Julia_2026_10_09'
def read(name):return json.loads((HERE/name).read_text(encoding='utf-8-sig'))
spec=read('lab_content.json');notes=read('teacher_notes.json');snippets=read('code_snippets.json');frames=read('animation_frames.json')
student_sources=[s for s in spec['sources'] if '/.codex/skills/' not in s['path_or_url']]
assert len(spec['objectives'])>=5 and len(spec['quiz'])==5
assert sum(x['minutes'] for x in spec['timeline'])==spec['duration']==180
assert sum(x['points'] for x in spec['rubric'])==100
E=lambda s:html.escape(str(s))
def p(s):return {'kind':'p','text':s}
def h(s,level=2):return {'kind':'h','text':s,'level':level}
def ul(values):return {'kind':'ul','values':values}
def code(text):return {'kind':'code','text':text}
def table(headers,rows):return {'kind':'table','headers':headers,'rows':rows}
book=[h('BÀI THỰC HÀNH CẤU TRÚC DỮ LIỆU VÀ GIẢI THUẬT',1),h(spec['title'],1),p('Ngày 09/10/2026 · Julia · Kịch bản đề xuất 180 phút'),p(spec['scenario']),h('Mục tiêu'),ul(spec['objectives']),h('Kịch bản thực hành'),table(['Thời lượng','Trọng tâm','Hoạt động','Minh chứng'],[[str(x['minutes'])+' phút',x['title'],x['activity'],x['evidence']] for x in spec['timeline']]),h('Chuẩn bị'),ul(spec['prerequisites'])]
book += [p('Ứng dụng xuyên suốt: hai CSV → danh mục đơn và bản đồ → thời gian từ BFS → ưu tiên và chọn đơn → hàng đợi → giao/hoàn tác → CSV kết quả. Bốn hàm cần cài đặt nằm trong cùng một file Julia; các cấu trúc và bộ điều phối được cung cấp.'),table(['Đơn','Điểm giao','Phục vụ','Lợi ích'],[['A / 1',2,1,6],['B / 2',4,1,8],['C / 3',6,1,10]]),p('Ngân sách của bộ mini là 8. Bản đồ vô hướng có sáu điểm, bảy cạnh; mỗi cạnh mất một đơn vị thời gian. Mỗi đơn là một chuyến riêng từ kho 1 tới điểm giao rồi trở về 1. Thời gian một đơn = 2 × số cạnh đường BFS + thời gian phục vụ.')]
setup=[h('Khởi động chương trình bằng Julia'),p('Giải nén bộ sinh viên; mở Julia tại thư mục chứa các file. Khung chạy được phần dữ liệu ngay. Những lời gọi cần thuật toán chưa hoàn thiện sẽ báo CHUA_HOAN_THANH.'),code('julia TH_2026_10_09_Sinh_vien.jl --data\njulia TH_2026_10_09_Kiem_thu.jl --phase=data'),p('Trong cửa sổ Julia, nạp khung và xem dữ liệu bằng các lệnh sau. Đây là phần đã được cung cấp, không cần cài gói CSV ngoài.'),code(snippets['setup']),code(snippets['read_csv']),p('Hai CSV quy định chỉ có số nguyên và tiêu đề cố định. Hàm đọc kiểm tra định dạng và giá trị; không áp dụng cách tách bằng dấu phẩy cho CSV tổng quát có ô chứa dấu phẩy hoặc xuống dòng.')]
setup += [p('Sau khi lưu phần code vừa sửa, khởi động lại Julia REPL và chạy lại các lệnh nạp khung, đọc dữ liệu ở trên trước khi thử thuật toán. Việc lưu file không tự cập nhật các hàm trong phiên Julia đang mở. Các lệnh julia ... --phase=... trong terminal mở một phiên mới cho mỗi lần kiểm tra.'),p('Cách chạy tệp và include: Julia Manual — Getting Started, https://docs.julialang.org/en/v1/manual/getting-started/. Phép đọc/ghi: Julia Base — I/O and Network, https://docs.julialang.org/en/v1/base/io-network/.')]
book+=setup
todo={1:('merge_sort','sort_call','sort'),2:('binary_search','search_call','search'),3:('bfs_routes','bfs_call','bfs'),4:('knapsack','dp_call','dp')}
def section_blocks(s):
 b=[h(s['title']),h('Lý thuyết sử dụng',3),ul(s['theory']),h('Ví dụ và suy luận',3),ul(s['example']),h('Lỗi cần tránh',3),ul(s['mistakes']),h('Chi phí và điều kiện áp dụng',3),ul(s['complexity'])]
 t=s.get('task')
 if t:
  n=int(re.findall(r'\d+',t['id'])[-1]);function,call,phase=todo[n]
  b += [h(t['id']+' · '+t['title'],3),ul(t['requirements']),h('Gợi ý triển khai',3),ul(t['hints']),code(snippets['signatures'][function]+'\n    # Sinh viên cài đặt tại vị trí TODO trong file chung.\nend'),h('Lời gọi kiểm tra trong cùng chương trình',3),code(snippets[call]),code('julia TH_2026_10_09_Kiem_thu.jl --phase='+phase),ul(t['verification'])]
 return b
for s in spec['sections']:book+=section_blocks(s)
closing=[h('Chạy ứng dụng sau khi hoàn thiện bốn hàm'),code('julia TH_2026_10_09_Sinh_vien.jl --mini\njulia TH_2026_10_09_Kiem_thu.jl --phase=app\njulia TH_2026_10_09_Kiem_thu.jl --phase=all'),p('Phần demo tính kế hoạch, giao một đơn, hoàn tác rồi giao hết kế hoạch. Theo dõi tuyến đường, hàng đợi, ngăn xếp, thời gian và lợi ích sau từng thao tác. Kết quả được xuất từ trạng thái thực tế, không nhập tay để thay đầu ra.'),code(snippets['application_call']),h('Nộp bài và trình bày'),ul(spec['submission']),p('Mẫu phiếu kèm theo yêu cầu ghi ý tưởng, bảng truy vết và kết quả kiểm thử. Trình bày bằng notebook Julia/Pluto gọi cùng lõi. Đưa toàn bộ code, bộ sinh/đọc CSV, kiểm thử và notebook lên GitHub, ghi URL cùng commit SHA.'),h('Tiêu chí đánh giá đề xuất'),table(['Tiêu chí','Điểm','Minh chứng'],[[r['criterion'],r['points'],r['evidence']] for r in spec['rubric']]),h('Câu hỏi cuối buổi')]
report_blocks=[h('Viết báo cáo theo cùng luồng ứng dụng'),p('Độ dài gợi ý 3–5 trang. Mỗi phần dùng dữ liệu và kết quả thực tế, dẫn tới hàm cùng commit trên GitHub.')]
for r in spec['report']:report_blocks += [h(str(r['number'])+'. '+r['title'],3),ul(r['content'])]
at=next(i for i,b in enumerate(closing) if b['kind']=='h' and b['text']=='Tiêu chí đánh giá đề xuất');closing[at:at]=report_blocks
for q in spec['quiz']:closing += [p(str(q['id'])+'. '+q['question'])]
closing += [h('Kiến thức bảy chương trong bài thực hành'),table(['Chương','Kiến thức','Vai trò trong luồng'],[[x['chapter'],x['knowledge'],x['role']] for x in spec['coverage']]),p('Trọng tâm cài đặt trong buổi học là bốn hàm đã nêu. Danh sách, hàng đợi, ngăn xếp, cây và các phép kiểm chứng được đọc và vận dụng qua khung ứng dụng; phần mở rộng được thực hiện sau buổi học.'),h('Nguồn học tập')]
for s in student_sources:closing += [p(s['title']+(' — '+s['path_or_url'] if s['path_or_url'].startswith('https://') else '')+'. '+s['note'])]
book+=closing

def to_html(blocks):
 out=[]
 for b in blocks:
  if b['kind']=='h':out.append('<h'+str(b['level'])+'>'+E(b['text'])+'</h'+str(b['level'])+'>')
  elif b['kind']=='p':out.append('<p>'+E(b['text'])+'</p>')
  elif b['kind']=='ul':out.append('<ul>'+''.join('<li>'+E(v)+'</li>' for v in b['values'])+'</ul>')
  elif b['kind']=='code':out.append('<pre><code>'+E(b['text'])+'</code></pre>')
  elif b['kind']=='table':out.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th scope="col">'+E(v)+'</th>' for v in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+E(v)+'</td>' for v in row)+'</tr>' for row in b['rows'])+'</tbody></table></div>')
 return ''.join(out)

def word(blocks,path,image_path=None):
 doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
 for name in ['Normal','Heading 1','Heading 2','Heading 3','List Bullet']:
  st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14);st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(5);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
  if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True
 st=doc.styles.add_style('Julia Code',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Consolas';st.font.size=Pt(10);st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_after=Pt(0)
 st=doc.styles.add_style('Compact Table',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Times New Roman';st.font.size=Pt(11);st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(3)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('CTDL & GT · Thực hành 09/10/2026 · Trang ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 for i,b in enumerate(blocks):
  if b['kind']=='h':doc.add_heading(b['text'],b['level'])
  elif b['kind']=='p':doc.add_paragraph(b['text'])
  elif b['kind']=='ul':
   for text in b['values']:doc.add_paragraph(text,'List Bullet')
  elif b['kind']=='code':
   for line in b['text'].splitlines():doc.add_paragraph(line or ' ','Julia Code')
  elif b['kind']=='table':
   t=doc.add_table(rows=1,cols=len(b['headers']));t.style='Table Grid'
   for c,v in zip(t.rows[0].cells,b['headers']):c.text=str(v)
   repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
   for row in b['rows']:
    for c,v in zip(t.add_row().cells,row):c.text=str(v)
   if len(b['headers'])>=5:
    for row in t.rows:
     for cell in row.cells:
      for para in cell.paragraphs:para.style='Compact Table'
  if i==3 and image_path and image_path.exists():
   inline=doc.add_picture(str(image_path),width=Cm(16.3));inline._inline.docPr.set('descr','Luồng xử lý: CSV, danh mục, BFS, chọn đơn, giao và hoàn tác, xuất CSV.')
 doc.save(path)

worksheet=[h('PHIẾU THỰC HÀNH — 09/10/2026',1),p('Họ tên: ........................................ MSSV: ........................ Lớp: ...................'),p('GitHub: ................................................ Commit SHA: ......................................'),h('1. Mô tả một luồng ứng dụng'),p('Ghi mục tiêu, đầu vào, đầu ra và sơ đồ truyền dữ liệu giữa các chức năng.'),p('...................................................................................................................'),h('2. Bảng truy vết đường và thời gian'),table(['ID','Đường lượt đi','Số cạnh h','2h + phục vụ'],[[1,'...................','........','........'],[2,'...................','........','........'],[3,'...................','........','........']]),h('3. Minh chứng bốn hàm')]
for s in spec['sections']:
 t=s.get('task')
 if t:worksheet += [h(t['id']+' · '+t['title'],3),p('Ý tưởng/điều kiện áp dụng: ........................................................................................'),p('Lời gọi và kết quả kiểm thử: ....................................................................................'),p('Chi phí thời gian/bộ nhớ: ........................................................................................')]
worksheet += [h('4. Quyết định chọn đơn'),p('Điền ngân sách, tập chọn, tổng thời gian/lợi ích. Giải thích một ô DP và một bước truy vết; đối chiếu với vét cạn.'),p('...................................................................................................................'),h('5. Giao và hoàn tác'),table(['Trạng thái','Queue','Stack','Đã dùng','Đã nhận'],[['Sau lập kế hoạch','........','........','........','........'],['Sau giao một đơn','........','........','........','........'],['Sau hoàn tác','........','........','........','........']]),h('6. Đầu ra và nhận xét'),p('Đường dẫn CSV kết quả: .................................... Link notebook: ................................'),p('Lỗi đã phát hiện và cách sửa: ..................................................................................'),p('Một chức năng thay đổi ảnh hưởng đến chức năng sau như thế nào? ................................'),h('Câu hỏi cuối buổi')]
for q in spec['quiz']:worksheet += [p(str(q['id'])+'. '+q['question']),p('...................................................................................................................')]

teacher=[h('HƯỚNG DẪN GIẢNG VIÊN — 09/10/2026',1),p('Tài liệu đáp án và điều hành buổi thực hành; phát bộ sinh viên riêng.'),p('Mạch ứng dụng và dữ liệu trùng bài thực hành chính. Không yêu cầu sinh viên cài lại tất cả cấu trúc dữ liệu trong một buổi.')]
labels={'title':'Tên tài liệu','date':'Ngày thực hành','audience':'Đối tượng sử dụng','defaultDuration':'Thời lượng đề xuất','teachingNotes':'Lưu ý điều hành','answers':'Đáp án câu hỏi cuối buổi','miniDataset':'Dữ liệu và kết quả bộ mini','dpExplanation':'Giải thích bảng quy hoạch động','integratedTrace':'Truy vết ứng dụng','changedMapExample':'Ví dụ đổi bản đồ','greedyCounterexample':'Phản ví dụ của tham lam','timeline90':'Phương án 90 phút','timeline225':'Phương án 225 phút','gradingGuidance':'Hướng dẫn chấm','oralQuestions':'Câu hỏi trao đổi','teacherCommands':'Lệnh chạy dành cho giảng viên','releaseChecklist':'Kiểm tra tài liệu trước buổi học','id':'Mã','question':'Câu hỏi','answer':'Đáp án','reason':'Giải thích','minutes':'Số phút','activity':'Hoạt động','evidence':'Minh chứng','budget':'Ngân sách','orders':'Đơn hàng','edges':'Cạnh bản đồ','costs':'Chi phí','chosen':'Tập chọn','totalCost':'Tổng chi phí','totalBenefit':'Tổng lợi ích','table':'Bảng','formula':'Công thức','trace':'Truy vết','criterion':'Tiêu chí','points':'Điểm'}
labels.update({'step':'Bước','selected':'Tập chọn','pending':'Hàng đợi','completed':'Lịch sử','spent':'Thời gian đã dùng','earned':'Lợi ích đã nhận','definition':'Định nghĩa','budgets':'Ngân sách','rows':'Các hàng','keyCells':'Các ô cần giải thích','gradingNote':'Lưu ý chấm','vertices':'Số đỉnh','destination':'Điểm giao','service':'Phục vụ','benefit':'Lợi ích','hops':'Số cạnh','cost':'Chi phí','route':'Tuyến','b':'Ngân sách','decision':'Quyết định','i':'Số đơn đã xét','label':'Nhãn','values':'Giá trị'})
def cell_text(value):return str(value) if not isinstance(value,(list,dict)) else json.dumps(value,ensure_ascii=False)
def flatten(value,label='',level=2):
 if isinstance(value,dict):
  out=[h(label,level)] if label else []
  if 'budgets' in value and 'rows' in value:
   out += [p(value['definition']),table(['i / b']+value['budgets'],[[row['i']]+row['values'] for row in value['rows']]),ul(value['keyCells']),table(['i','b','Quyết định'],[[t['i'],t['b'],t['decision']] for t in value['trace']]),p(value['gradingNote'])]
   return out
  for k,v in value.items():out+=flatten(v,labels.get(k,k),min(level+1,3))
  return out
 if isinstance(value,list):
  out=[h(label,level)] if label else []
  if value and all(isinstance(x,dict) and x.keys()==value[0].keys() and all(not isinstance(y,dict) for y in x.values()) for x in value):
   keys=list(value[0]);return out+[table([labels.get(k,k) for k in keys],[[cell_text(x[k]) for k in keys] for x in value])]
  if all(not isinstance(x,(dict,list)) for x in value):return out+[ul([str(x) for x in value])]
  for v in value:out+=flatten(v,'',level)
  return out
 return [p((label+': ' if label else '')+str(value))]
teacher+=flatten(notes)
teacher += [h('Code tham khảo bốn hàm trọng tâm')]
for name,text in snippets['functions'].items():teacher += [h(name,3),code(text)]
teacher += [h('Chạy lời giải và kiểm thử'),code('julia TH_2026_10_09_Loi_giai.jl --mini\njulia TH_2026_10_09_Kiem_thu.jl --teacher --phase=all')]

for source,target in [('lab_student.jl','TH_2026_10_09_Sinh_vien.jl'),('lab_solution.jl','TH_2026_10_09_Loi_giai.jl'),('test_lab.jl','TH_2026_10_09_Kiem_thu.jl')]:shutil.copyfile(HERE/source,ROOT/target)
shutil.copyfile(HERE/'notebook_template.jl',ROOT/'TH_2026_10_09_Pluto.jl')
word(book,ROOT/(BASE+'.docx'),HERE/'qa/roadmap.png');word(worksheet,ROOT/'Phieu_thuc_hanh_CTDL_GT_2026_10_09.docx');word(teacher,ROOT/'Huong_dan_giang_vien_CTDL_GT_2026_10_09.docx')
css=(HERE/'lesson.css').read_text(encoding='utf8')+(HERE/'animations.css').read_text(encoding='utf8')
roadmap='<div id="roadmap" class="roadmap" role="img" aria-label="Các bước của cùng ứng dụng giao đơn">'+''.join('<article><strong>'+E(a)+'</strong><span>'+E(b)+'</span></article>' for a,b in [('CSV','Đơn và bản đồ'),('Danh mục','Sắp xếp, tra cứu'),('BFS','Đường → thời gian'),('DP','Chọn trong B'),('Giao / hoàn tác','Queue / stack'),('CSV kết quả','Cùng trạng thái')])+'</div>'
sections=[]
def section(id,content):sections.append('<section id="'+id+'" class="lesson-section">'+content+'</section>')
section('gioi-thieu','<p class="eyebrow">HỌC PHẦN INT1306 · THỰC HÀNH 09/10/2026</p><h1>'+E(spec['title'])+'</h1><p class="lead">Một bài toán, một chương trình Julia, bốn hàm trọng tâm.</p><p>'+E(spec['scenario'])+'</p>'+roadmap+'<p class="callout">Khung thời lượng: 180 phút. Giảng viên có hướng dẫn rút gọn 90 phút và mở rộng 225 phút trong tài liệu riêng.</p><div class="actions"><a class="button" href="Bo_Thuc_hanh_Julia_2026_10_09_Sinh_vien.zip">Tải bộ sinh viên</a><button id="download-code">Tải một file Julia mẫu</button><a class="button" href="'+BASE+'.pdf">Bài thực hành PDF</a><a class="button" href="'+BASE+'.docx">Bài thực hành Word</a></div>')
section('muc-tieu',to_html([h('Mục tiêu và tổ chức buổi học'),ul(spec['objectives']),table(['Phút','Trọng tâm','Hoạt động','Kết quả'],[[x['minutes'],x['title'],x['activity'],x['evidence']] for x in spec['timeline']])]))
section('bai-toan',to_html([h('Một bài toán dùng chung dữ liệu'),p('Ba đơn A, B, C dùng chung bản đồ. Mỗi chuyến đi và về kho 1; ngân sách B=8.'),table(['ID','Điểm giao','Phục vụ','Lợi ích'],[[1,2,1,6],[2,4,1,8],[3,6,1,10]]),p('Đọc đơn → tra cứu và tạo thứ tự → BFS tính thời gian → DP chọn đơn → đưa kế hoạch vào hàng đợi → giao/hoàn tác → xuất CSV.'),p('Nếu đường thay đổi, thời gian có thể thay đổi; kế hoạch và hàng đợi sử dụng thời gian mới. Bài tập yêu cầu bảo toàn mối liên hệ này.')]))
section('khoi-dong',to_html([h('Chuẩn bị trước khi cài thuật toán'),ul(spec['prerequisites'])]+setup))
section('minh-hoa','<h2>Minh họa động từ cùng bộ dữ liệu</h2><p>Chọn một thuật toán, quan sát từng bước và ghi trạng thái trước khi viết Julia. Khung này minh họa thuật toán; chương trình thực hành chạy bằng Julia.</p><div id="lab-animation"></div>')
for s in spec['sections']:section(s['id'],to_html(section_blocks(s)))
section('tich-hop-nop-bai',to_html(closing[:closing.index(next(b for b in closing if b['kind']=='h' and b['text']=='Câu hỏi cuối buổi'))]))
section('cau-hoi','<h2>Câu hỏi cuối buổi</h2><ol class="quiz">'+''.join('<li>'+E(q['question'])+'<div class="blank"></div></li>' for q in spec['quiz'])+'</ol><div class="actions"><a class="button" href="Phieu_thuc_hanh_CTDL_GT_2026_10_09.docx">Phiếu thực hành Word</a><a class="button" href="Phieu_thuc_hanh_CTDL_GT_2026_10_09.pdf">Phiếu thực hành PDF</a></div>')
sources=[]
for s in student_sources:
 ref=s['path_or_url'];label='<a href="'+E(ref)+'" target="_blank" rel="noopener">'+E(s['title'])+'</a>' if ref.startswith('https://') else E(s['title'])
 sources.append('<li>'+label+'<br>'+E(s['note'])+'</li>')
section('phu-luc',to_html([h('Vị trí trong chương trình học'),table(['Chương','Kiến thức','Vai trò'],[[x['chapter'],x['knowledge'],x['role']] for x in spec['coverage']]),p('Bốn hàm là phần cài đặt trong lớp. Các cấu trúc đã cung cấp được đọc, truy vết và vận dụng trong cùng ứng dụng.')])+'<h3>Nguồn học tập</h3><ul class="source-list">'+''.join(sources)+'</ul>')
safe=lambda x:json.dumps(x,ensure_ascii=False).replace('</','<\\/')
head='<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Thực hành CTDLGT · Julia · 09/10/2026</title><style>'+css+'</style></head><body>'
page=head+'<header class="top"><strong>CTDL & GT · Thực hành Julia</strong><nav><a href="#khoi-dong">Bắt đầu</a><a href="#minh-hoa">Minh họa</a><a href="#cau-hoi">Phiếu bài tập</a><button id="projection" aria-pressed="false">Trình chiếu</button><button id="font-plus">Đổi cỡ chữ</button><button id="print-lesson">In / lưu PDF</button></nav></header><main>'+''.join(sections)+'</main><nav class="slide-bar" aria-label="Điều khiển trình chiếu"><button id="slide-prev">← Trước</button><select id="slide-picker" aria-label="Chọn trang"></select><span id="slide-number"></span><button id="slide-next">Tiếp →</button></nav><footer>Bài thực hành 09/10/2026 · Một ứng dụng Julia · Bốn hàm trọng tâm · Dữ liệu CSV giả lập.</footer><script>window.LabFrames='+safe(frames)+';window.LabResources='+safe({'student':(HERE/'lab_student.jl').read_bytes().decode('utf8')})+';</script><script>'+(HERE/'animations.js').read_text(encoding='utf8')+'</script><script>'+(HERE/'lesson.js').read_text(encoding='utf8')+'</script></body></html>'
(ROOT/(BASE+'.html')).write_text(page,encoding='utf8')
preview_css=(HERE/'lesson.css').read_text(encoding='utf8')+'body{background:white}main{max-width:1050px}.lesson-section{border:0} @media print{.lesson-section{break-before:auto}h2{margin-top:18pt}}'
for name,blocks in [('worksheet_preview',worksheet),('teacher_preview',teacher)]:
 (HERE/(name+'.html')).write_text('<!doctype html><html lang="vi"><meta charset="utf-8"><style>'+preview_css+'</style><main><section class="lesson-section">'+to_html(blocks)+'</section></main></html>',encoding='utf8')
md=[]
for b in book:
 if b['kind']=='h':md.append('#'*b['level']+' '+b['text'])
 elif b['kind']=='p':md.append(b['text'])
 elif b['kind']=='ul':md.append('\n'.join('- '+str(x) for x in b['values']))
 elif b['kind']=='code':md.append('```julia\n'+b['text']+'\n```')
 elif b['kind']=='table':md.append('| '+' | '.join(map(str,b['headers']))+' |\n| '+' | '.join('---' for _ in b['headers'])+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in b['rows']))
(ROOT/(BASE+'.md')).write_text('\n\n'.join(md),encoding='utf8')
print(json.dumps({'status':'BUILD_PASS','date':'2026-10-09','duration':180,'tasks':4,'slides':len(sections),'quiz':5}))
